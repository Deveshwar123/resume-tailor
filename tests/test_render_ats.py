"""Milestone 6: rendering, page-fit order, and ATS checks."""

from __future__ import annotations

import pytest

import config
from src import ats_check, render, tailor_resume, validate
from src.models import BulletFix, BulletFixBatch
from src.relevance import Relevance
from src.render import FitResult, LayoutParams, ResumeView, ViewEntry, ViewSection
from tests.conftest import ScriptedLLM, bd, ed, make_draft

needs_soffice = pytest.mark.skipif(render.find_soffice() is None, reason="LibreOffice (soffice) is not installed")


@pytest.fixture
def resume(ctx):
    draft = make_draft(
        experiences=[
            ed("exp_lab", [bd("Participated in weekly lab meetings and annotated 500 images", ["exp_lab_f01"], None, "unchanged")]),
            ed("exp_acme", [bd("Assisted a 4-person team in developing a Python application.", ["exp_acme_f01"], "exp_acme_b01", "unchanged"),
                            bd("Wrote SQL queries that reduced report generation time by 30%.", ["exp_acme_f02"], "exp_acme_b02", "unchanged")]),
            ed("exp_club", [bd("Led weekly workshops for 25 members on data analysis", ["exp_club_f01"], None, "unchanged"),
                            bd("Organized a campus hackathon with 60+ participants", ["exp_club_f02"], None, "unchanged")]),
        ],
        projects=[ed("proj_tracker", [bd("Built a budget tracking web app with JavaScript and React.", ["proj_tracker_f01"],
                                         "proj_tracker_b01", "unchanged"),
                                      bd("Deployed the app for 40 student users", ["proj_tracker_f02"], None, "unchanged")])],
        skills=["Python", "SQL", "React", "data analysis"], honors=["Dean's List"], coursework=["Databases"],
        award_ids=["award_hack"], order=["education", "experience", "projects", "leadership", "skills", "awards"],
    )
    return validate.enforce(draft, ctx).resume


# ---------------------------------------------------------------------------
# View
# ---------------------------------------------------------------------------


def test_view_copies_headers_from_the_master_resume(resume, master):
    view = render.build_view(resume, master)
    assert view.name == "TEST PERSON"
    assert view.contact_line == "Austin, TX | (555) 010-0100 | test.person@example.com | linkedin.com/in/testperson | github.com/testperson"
    experience = view.section("experience")
    # Reverse chronological regardless of the order the AI chose (lab was listed first).
    assert [e.source_id for e in experience.entries] == ["exp_acme", "exp_lab"]
    acme = experience.entries[0]
    assert acme.left == "Acme Corp — Software Engineering Intern, Austin, TX"
    assert acme.right == "Jun – Aug 2025"
    assert all(b.endswith(".") for b in acme.bullets)
    assert [e.source_id for e in view.section("leadership").entries] == ["exp_club"]     # clubs go under LEADERSHIP
    assert view.section("projects").entries[0].left == "Budget Tracker"


def test_view_education_skills_and_awards(resume, master):
    view = render.build_view(resume, master)
    edu = view.section("education").entries[0]
    assert edu.left == "State University, Austin, TX" and edu.right == "May 2028"
    assert edu.lines == ["B.S. Computer Science", "Cumulative GPA (GPR): 3.6/4.0", "Honors: Dean's List",
                         "Relevant Coursework: Databases"]
    assert view.section("skills").lines == ["Languages: Python, SQL", "Tools and Platforms: React", "Additional Skills: data analysis"]
    assert view.section("awards").bullets == ["Campus Hackathon winner, 2025."]
    assert [s.heading for s in view.sections] == ["EDUCATION", "EXPERIENCE", "PROJECTS", "LEADERSHIP", "SKILLS",
                                                  "AWARDS AND RECOGNITION"]


def test_plain_text_is_what_the_judge_reads(resume, master):
    text = render.build_view(resume, master).plain_text()
    assert "EXPERIENCE\nAcme Corp — Software Engineering Intern, Austin, TX  Jun – Aug 2025\n- Assisted a 4-person team" in text


# ---------------------------------------------------------------------------
# DOCX / PDF
# ---------------------------------------------------------------------------


def test_docx_is_filled_from_the_template(resume, master, tmp_path):
    from docx import Document

    view = render.build_view(resume, master)
    path = render.render_docx(view, tmp_path / "resume.docx")
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "TEST PERSON" and "EXPERIENCE" in texts
    assert "Acme Corp — Software Engineering Intern, Austin, TX\tJun – Aug 2025" in texts
    assert "•\tAssisted a 4-person team in developing a Python application." in texts
    assert not any("{{" in t or "{%" in t for t in texts)           # no template tags left behind
    assert ats_check.check_docx(path, view.atoms()).passed


def test_layout_changes_stay_inside_the_document(resume, master, tmp_path):
    from docx import Document

    view = render.build_view(resume, master)
    path = render.render_docx(view, tmp_path / "r.docx", LayoutParams(margin_in=0.5, font_scale=10 / 11, spacing_scale=0.4))
    doc = Document(str(path))
    assert doc.sections[0].left_margin.inches == pytest.approx(0.5)
    assert doc.styles["Normal"].font.size.pt == pytest.approx(10.0)
    assert doc.styles["Resume Heading"].paragraph_format.space_before.pt == pytest.approx(7 * 0.4, abs=0.1)
    right_tab = [t for t in doc.styles["Resume Entry"].paragraph_format.tab_stops][0]
    assert right_tab.position.inches == pytest.approx(7.5, abs=0.01)   # dates stay on the right margin


def test_template_info_reads_the_template():
    margin, font = render.template_info()
    assert margin == pytest.approx(config.DEFAULT_MARGIN_IN) and font == pytest.approx(config.DEFAULT_FONT_PT)


@needs_soffice
def test_full_render_passes_every_ats_check(resume, master, analysis, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "CACHE_DIR", tmp_path / "cache")
    rendered = render.render(resume, master, tmp_path / "out")
    assert rendered.pages == 1 and rendered.pdf.exists() and rendered.docx.exists()
    report = ats_check.check_rendered(rendered.view, rendered.pdf, rendered.docx, analysis)
    assert report.final and [c.name for c in report.failures()] == []
    assert report.score == pytest.approx(10.0)
    assert {c.name for c in report.checks} >= {"contact_info", "text_extraction", "section_headings", "odd_characters",
                                                "links", "keyword_repetition", "reverse_chronological", "page_count"}


# ---------------------------------------------------------------------------
# ATS: text-extraction comparison
# ---------------------------------------------------------------------------

ATOMS = ["TEST PERSON", "EDUCATION", "State University, Austin, TX", "May 2028", "EXPERIENCE",
         "Acme Corp — Software Engineering Intern", "Wrote SQL queries that reduced report generation time by 30%.",
         "SKILLS", "Languages: Python, SQL"]


def test_extraction_matches_despite_wraps_and_bullets():
    extracted = ("TEST PERSON\nEDUCATION\nState University, Austin, TX May 2028\nEXPERIENCE\n"
                 "Acme Corp — Software Engineering Intern\n• Wrote SQL queries that reduced report \ngeneration time by 30%.\n"
                 "SKILLS\nLanguages: Python, SQL")
    check = ats_check.compare_extraction(ATOMS, extracted)
    assert check.passed and check.score == 1.0 and "100.0%" in check.detail


def test_extraction_reports_a_missing_section():
    extracted = "TEST PERSON\nEDUCATION\nState University, Austin, TX May 2028\nSKILLS\nLanguages: Python, SQL"
    check = ats_check.compare_extraction(ATOMS, extracted)
    assert not check.passed and "missing" in check.detail and check.score < 0.7


def test_extraction_reports_garbled_characters():
    extracted = "\n".join(ATOMS).replace("queries", "que�ries") + "\n(cid:12)"
    check = ats_check.compare_extraction(ATOMS, extracted)
    assert not check.passed and "garbled" in check.detail


def test_extraction_reports_wrong_reading_order():
    lines = list(ATOMS)
    lines[5], lines[6] = lines[6], lines[5]          # a two-column layout would interleave like this
    scrambled = "\n".join(reversed(lines))
    check = ats_check.compare_extraction(ATOMS, scrambled)
    assert not check.passed and "out of reading order" in check.detail


# ---------------------------------------------------------------------------
# ATS: the other checks
# ---------------------------------------------------------------------------


def _view(**kw) -> ResumeView:
    base = dict(name="TEST PERSON", contact_line="Austin, TX | (555) 010-0100 | t@example.com",
                contact={"name": "Test Person", "email": "t@example.com", "phone": "(555) 010-0100",
                         "linkedin": "linkedin.com/in/test", "github": ""},
                sections=[ViewSection("education", "EDUCATION", lines=["B.S."]),
                          ViewSection("experience", "EXPERIENCE", entries=[
                              ViewEntry("Acme — Intern", "Jun – Aug 2025", bullets=["Wrote Python."], start="2025-06", end="2025-08"),
                              ViewEntry("Old Co — Intern", "2023", bullets=["Wrote Java."], start="2023", end="2023")]),
                          ViewSection("skills", "SKILLS", lines=["Languages: Python, Java"])])
    base.update(kw)
    return ResumeView(**base)


def test_text_level_report_is_clean_for_a_good_view():
    report = ats_check.check_text(_view())
    assert not report.final and report.failures() == [] and report.score == 10.0


def test_contact_problems_are_caught():
    view = _view(contact={"name": "", "email": "not-an-email", "phone": "555", "linkedin": "", "github": ""})
    check = ats_check.check_contact(view)
    assert not check.passed and "name" in check.detail and "email" in check.detail and "phone" in check.detail


def test_missing_required_heading_is_caught():
    view = _view()
    view.sections = [s for s in view.sections if s.name != "experience"]
    assert 'missing the "EXPERIENCE" section' in ats_check.check_headings(view).detail


def test_odd_characters_are_caught():
    view = _view()
    view.sections[1].entries[0].bullets = ["Shipped \U0001F680 features ➢ fast with ﬁne tools"]
    check = ats_check.check_odd_characters(view)
    assert not check.passed and "U+1F680" in check.detail and "U+FB01" in check.detail
    assert ats_check.odd_characters("Café résumé • 2024–2025 — fine") == []


def test_malformed_links_are_caught():
    view = _view(contact={"name": "T", "email": "t@example.com", "phone": "(555) 010-0100",
                          "linkedin": "linkedin dot com/in/test", "github": "github.com/test"})
    assert not ats_check.check_links(view).passed
    assert ats_check.check_links(_view()).passed


def test_keyword_stuffing_is_flagged():
    view = _view()
    view.sections[1].entries[0].bullets = ["Python scripts.", "Python tests.", "Python tools and more Python."]
    check = ats_check.check_repetition(view, ["Python", "Java"])
    assert not check.passed and "Python (5x)" in check.detail      # 4 in bullets + 1 in the skills line


def test_out_of_order_entries_are_flagged():
    view = _view()
    view.sections[1].entries.reverse()
    assert "not in reverse chronological order" in ats_check.check_chronology(view).detail
    assert ats_check.check_chronology(_view()).passed


def test_page_count_must_be_exactly_one():
    assert ats_check.check_page_count(1).passed
    assert not ats_check.check_page_count(2).passed and not ats_check.check_page_count(0).passed


# ---------------------------------------------------------------------------
# Page fit: the order of steps
# ---------------------------------------------------------------------------

STEP_ORDER = ["render", "tighten", "spacing", "margins", "font", "remove_bullet", "remove_project",
              "remove_award", "remove_entry"]


def fit(resume, master, analysis, render_fn, **kw) -> FitResult:
    return render.fit_to_page(resume, master, render_fn, relevance=Relevance(analysis), **kw)


def shorten_all(items):
    return {i.key: i.text[: max(20, len(i.text) // 2)] for i in items}


def test_nothing_happens_when_it_already_fits(resume, master, analysis):
    result = fit(resume, master, analysis, lambda r, layout: 1, tighten_fn=shorten_all)
    assert result.fits and [s.kind for s in result.steps] == ["render"] and result.removed == []


def test_steps_run_in_the_spec_order_and_respect_the_limits(resume, master, analysis, monkeypatch):
    monkeypatch.setattr(config, "TIGHTEN_MIN_CHARS", 10)
    layouts: list[LayoutParams] = []

    def never_fits(r, layout):
        layouts.append(LayoutParams(layout.margin_in, layout.font_scale, layout.spacing_scale))
        return 2

    result = fit(resume, master, analysis, never_fits, tighten_fn=shorten_all)
    kinds = [s.kind for s in result.steps]
    assert not result.fits
    # Each kind of step appears, and never before an earlier kind has finished.
    assert [k for k in STEP_ORDER if k in kinds] == STEP_ORDER
    assert kinds == sorted(kinds, key=STEP_ORDER.index)
    # ATS-safe limits are never crossed.
    assert min(l.spacing_scale for l in layouts) == pytest.approx(config.MIN_SPACING_SCALE)
    assert min(l.margin_in for l in layouts if l.margin_in is not None) == pytest.approx(config.MIN_MARGIN_IN)
    assert min(l.font_scale for l in layouts) * config.DEFAULT_FONT_PT == pytest.approx(config.MIN_FONT_PT)
    # Layout is exhausted before any content is removed.
    first_removal = kinds.index("remove_bullet")
    assert layouts[first_removal - 1].margin_in == pytest.approx(config.MIN_MARGIN_IN)


def test_tightening_alone_can_fix_it(resume, master, analysis, monkeypatch):
    monkeypatch.setattr(config, "TIGHTEN_MIN_CHARS", 10)
    before = sum(len(b.text) for _e, _i, b in resume.bullets())
    result = fit(resume, master, analysis,
                 lambda r, layout: 1 if sum(len(b.text) for _e, _i, b in r.bullets()) < before else 2,
                 tighten_fn=shorten_all)
    assert result.fits and [s.kind for s in result.steps] == ["render", "tighten"]
    assert any("page-fit: tightened" in note for _e, _i, b in resume.bullets() for note in b.history)


def test_rejected_tightening_changes_nothing(resume, master, analysis, monkeypatch):
    monkeypatch.setattr(config, "TIGHTEN_MIN_CHARS", 10)
    before = [b.text for _e, _i, b in resume.bullets()]
    result = fit(resume, master, analysis, lambda r, layout: 1 if layout.spacing_scale < 1 else 2,
                 tighten_fn=shorten_all, accept_fn=lambda entry, bullet, text: False)
    assert [b.text for _e, _i, b in resume.bullets()] == before
    assert [s.kind for s in result.steps][:3] == ["render", "tighten", "spacing"] and result.fits


def test_spacing_then_margins_then_font(resume, master, analysis):
    result = fit(resume, master, analysis, lambda r, layout: 1 if layout.font_scale < 1 else 2)
    kinds = [s.kind for s in result.steps]
    assert result.fits and kinds[-1] == "font" and "remove_bullet" not in kinds
    assert kinds.index("spacing") < kinds.index("margins") < kinds.index("font")
    assert result.layout.margin_in == pytest.approx(config.MIN_MARGIN_IN)


def test_removal_takes_the_lowest_relevance_bullet_first_then_the_project(resume, master, analysis):
    # The posting wants Python and SQL for a data role. Several bullets mention none of its terms;
    # among those, the one in the least relevant entry (the web project) goes first.
    total = lambda r: sum(len(e.bullets) for e in r.entries())  # noqa: E731
    start = total(resume)
    result = fit(resume, master, analysis, lambda r, layout: 1 if total(r) <= start - 1 and layout.font_scale < 1 else 2)
    assert result.fits and result.steps[-1].kind == "remove_bullet"
    assert result.removed == ['bullet from proj_tracker: "Deployed the app for 40 student users"']

    result = fit(resume, master, analysis, lambda r, layout: 1 if not r.projects else 2)
    kinds = [s.kind for s in result.steps]
    assert result.fits and kinds[-1] == "remove_project" and kinds.index("remove_bullet") < kinds.index("remove_project")
    assert all(len(e.bullets) >= config.MIN_BULLETS_PER_ENTRY for e in resume.experiences)   # no entry is emptied


def test_tightened_text_is_revalidated(ctx, resume):
    accept = render.make_accept_fn(ctx)
    entry = next(e for e in resume.experiences if e.source_id == "exp_acme")
    bullet = entry.bullets[0]
    assert accept(entry, bullet, "Assisted a 4-person team developing a Python application.")
    assert not accept(entry, bullet, "Led a 4-person team developing a Python application.")     # ownership
    assert not accept(entry, bullet, "Assisted a team developing a Python application.")          # dropped metric
    assert not accept(entry, bullet, "Assisted a 5-person team developing a Python application.")  # changed number


def test_tighten_call_reuses_the_cached_system_prompt(ctx):
    llm = ScriptedLLM({"tighten": [BulletFixBatch(fixes=[BulletFix(key="exp_acme#0", rewritten_text="Short.", fact_ids=[])])]})
    system = tailor_resume.build_system(ctx)
    out = render.make_tighten_fn(llm, system)([render.TightenItem("exp_acme#0", "exp_acme", "A long bullet.", ["exp_acme_f01"])])
    assert out == {"exp_acme#0": "Short."} and llm.calls[0]["system"] == system and llm.calls[0]["model"] == config.WRITE_MODEL


def test_section_override_moves_a_role(master_data, status, analysis, resume):
    from src.models import MasterResume

    master_data["experience"][1]["section"] = "experience"        # the club role, listed under EXPERIENCE
    moved = MasterResume.model_validate(master_data)
    view = render.build_view(resume, moved)
    assert "exp_club" in [e.source_id for e in view.section("experience").entries]
    assert view.section("leadership") is None
