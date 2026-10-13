"""ATS compatibility checks (pure Python, no AI).

Two entry points:

* `check_text(view, ...)`      - the checks that need no rendered file. Used inside
                                 the revision loop as the ATS part of Resume Quality.
* `check_rendered(view, pdf, docx, ...)` - every check, on the final PDF and DOCX.

Checks: contact info present, PDF text extraction matches the intended content
(missing text, garbled characters, reading order), required section headings,
no odd characters, well-formed links, no keyword stuffing, reverse-chronological
entries, the DOCX structure (one column, no tables or text boxes), and exactly
one page. The weighted result is the ATS category of Resume Quality.
"""

from __future__ import annotations

import re
from pathlib import Path

import config
from src import textutil
from src.models import AtsCheck, AtsReport, JobAnalysis
from src.render import ResumeView

# Printable ASCII, newline/tab, the bullet, en/em dashes, and accented Latin letters.
_ODD_CHAR_RE = re.compile(r"[^\x20-\x7E\n\t•–—À-ÖØ-öø-ÿ]")
_GARBLED_RE = re.compile(r"�|[-]|\(cid:\d+\)")
_EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")
_LINK_RE = re.compile(r"^(https?://)?([A-Za-z0-9-]+\.)+[A-Za-z]{2,}(/[^\s]*)?$")
_ORDER_MIN_CHARS = 12       # shorter pieces ("2024", "SKILLS") are not used to judge reading order


# ---------------------------------------------------------------------------
# Individual checks
# ---------------------------------------------------------------------------


def odd_characters(text: str) -> list[str]:
    """Distinct characters an ATS may mangle (emoji, fancy bullets, ligatures, ...)."""
    seen: dict[str, None] = {}
    for char in _ODD_CHAR_RE.findall(text or ""):
        seen.setdefault(char)
    return list(seen)


def _describe_chars(chars: list[str]) -> str:
    return ", ".join(f"{c!r} (U+{ord(c):04X})" for c in chars[:6]) + (" ..." if len(chars) > 6 else "")


def check_contact(view: ResumeView, extracted: str | None = None) -> AtsCheck:
    problems = []
    name, email, phone = view.contact.get("name", ""), view.contact.get("email", ""), view.contact.get("phone", "")
    if not name.strip():
        problems.append("name is missing")
    if not _EMAIL_RE.match(email.strip()):
        problems.append("email is missing or malformed")
    if len(re.sub(r"\D", "", phone)) < 10:
        problems.append("phone is missing or too short")
    if extracted is not None:
        hay = textutil.squash(extracted)
        for label, value in (("name", name), ("email", email), ("phone", phone)):
            if value.strip() and textutil.squash(value) not in hay:
                problems.append(f"{label} does not survive PDF text extraction")
    return AtsCheck(name="contact_info", passed=not problems, score=0.0 if problems else 1.0,
                    detail="; ".join(problems) or "name, email, and phone present")


def check_headings(view: ResumeView, extracted: str | None = None) -> AtsCheck:
    present = {s.name for s in view.sections}
    problems = [f'missing the "{config.SECTION_HEADINGS[name]}" section' for name in config.REQUIRED_SECTIONS
                if name not in present]
    if extracted is not None:
        hay = textutil.squash(extracted)
        problems += [f'heading "{s.heading}" does not survive PDF text extraction' for s in view.sections
                     if textutil.squash(s.heading) not in hay]
    score = max(0.0, 1.0 - len(problems) / max(1, len(config.REQUIRED_SECTIONS)))
    return AtsCheck(name="section_headings", passed=not problems, score=score if problems else 1.0,
                    detail="; ".join(problems) or "standard headings: " + ", ".join(s.heading for s in view.sections))


def check_odd_characters(view: ResumeView, extracted: str | None = None) -> AtsCheck:
    found = odd_characters("\n".join(view.atoms()))
    where = "the resume text"
    if not found and extracted is not None:
        found, where = odd_characters(extracted), "the extracted PDF text"
    if found:
        return AtsCheck(name="odd_characters", passed=False, score=max(0.0, 1.0 - 0.25 * len(found)),
                        detail=f"unusual characters in {where}: {_describe_chars(found)}")
    return AtsCheck(name="odd_characters", passed=True, detail="only plain characters, standard bullets, and dashes")


def check_links(view: ResumeView, resolve: bool = False) -> AtsCheck:
    """Links must be well-formed. Optionally HEAD-check my own links."""
    links = [view.contact.get(k, "").strip() for k in ("linkedin", "github")]
    links = [link for link in links if link]
    bad = [link for link in links if not _LINK_RE.match(link)]
    notes: list[str] = []
    if resolve and not bad:
        import requests

        for link in links:
            url = link if link.startswith("http") else f"https://{link}"
            try:
                status = requests.head(url, timeout=8, allow_redirects=True,
                                       headers={"User-Agent": config.USER_AGENT}).status_code
            except requests.RequestException:
                notes.append(f"{link}: could not be reached")
                continue
            if status == 404:
                bad.append(link)
                notes.append(f"{link}: not found (404)")
            elif status >= 400:
                notes.append(f"{link}: could not be verified (HTTP {status})")   # LinkedIn blocks bots
    if bad:
        return AtsCheck(name="links", passed=False, score=0.0, detail="problem links: " + ", ".join(bad)
                        + ("; " + "; ".join(notes) if notes else ""))
    detail = f"{len(links)} link(s) well-formed" if links else "no links on the resume"
    return AtsCheck(name="links", passed=True, detail=detail + ("; " + "; ".join(notes) if notes else ""))


def check_repetition(view: ResumeView, terms: list[str]) -> AtsCheck:
    """Flag any keyword used more than KEYWORD_REPEAT_LIMIT times (keyword stuffing)."""
    body = view.body_text()
    counts: dict[str, tuple[str, int]] = {}
    for term in terms:
        key = textutil.canon(term)
        if key and key not in counts:
            counts[key] = (term, textutil.count_term(body, term))
    over = sorted(((t, n) for t, n in counts.values() if n > config.KEYWORD_REPEAT_LIMIT), key=lambda x: -x[1])
    if over:
        return AtsCheck(name="keyword_repetition", passed=False, score=max(0.0, 1.0 - 0.25 * len(over)),
                        detail="used more than %d times: %s" % (config.KEYWORD_REPEAT_LIMIT,
                                                                ", ".join(f"{t} ({n}x)" for t, n in over)))
    return AtsCheck(name="keyword_repetition", passed=True,
                    detail=f"no keyword appears more than {config.KEYWORD_REPEAT_LIMIT} times")


def check_chronology(view: ResumeView) -> AtsCheck:
    problems = []
    for section in view.sections:
        dated = [e for e in section.entries if e.start or e.end]
        keys = [(textutil.parse_ym(e.end or e.start, end=True), textutil.parse_ym(e.start)) for e in dated]
        if keys != sorted(keys, reverse=True):
            problems.append(f"{section.heading} is not in reverse chronological order")
    return AtsCheck(name="reverse_chronological", passed=not problems, score=0.0 if problems else 1.0,
                    detail="; ".join(problems) or "entries are most-recent-first in every section")


def check_page_count(pages: int) -> AtsCheck:
    ok = pages == config.MAX_PAGES
    return AtsCheck(name="page_count", passed=ok, score=1.0 if ok else 0.0,
                    detail=f"{pages} page(s)" + ("" if ok else f"; must be exactly {config.MAX_PAGES}"))


def compare_extraction(atoms: list[str], extracted: str) -> AtsCheck:
    """Does the text an ATS would extract from the PDF match what we meant to write?

    Every intended piece of text is looked up in the extracted text, in reading
    order. Reports missing content, garbled characters, and wrong reading order.
    """
    hay = textutil.squash(extracted)
    garbled = sorted(set(_GARBLED_RE.findall(extracted or "")))
    cursor, total, matched = 0, 0, 0
    missing: list[str] = []
    out_of_order: list[str] = []
    for atom in atoms:
        needle = textutil.squash(atom)
        if not needle:
            continue
        total += len(needle)
        if len(needle) < _ORDER_MIN_CHARS:
            if needle in hay:
                matched += len(needle)
            else:
                missing.append(atom)
            continue
        position = hay.find(needle, cursor)
        if position >= 0:
            matched += len(needle)
            cursor = position + len(needle)
        elif needle in hay:
            matched += len(needle)
            out_of_order.append(atom)
        else:
            missing.append(atom)
    ratio = matched / total if total else 0.0
    problems = []
    if missing:
        problems.append(f"{len(missing)} piece(s) of text missing, e.g. \"{missing[0][:60]}\"")
    if out_of_order:
        problems.append(f"{len(out_of_order)} piece(s) out of reading order, e.g. \"{out_of_order[0][:60]}\"")
    if garbled:
        problems.append("garbled characters: " + ", ".join(repr(g) for g in garbled[:4]))
    passed = ratio >= config.EXTRACTION_MATCH_THRESHOLD and not out_of_order and not garbled
    score = ratio * (0.5 if out_of_order else 1.0) * (0.5 if garbled else 1.0)
    detail = f"{ratio:.1%} of the intended text extracted" + ("; " + "; ".join(problems) if problems else " in the right order")
    return AtsCheck(name="text_extraction", passed=passed, score=1.0 if passed and ratio >= 0.999 else score, detail=detail)


def check_docx(docx: Path, atoms: list[str]) -> AtsCheck:
    """The DOCX must be one column with no tables, text boxes, images, or header/footer text."""
    from docx import Document

    doc = Document(str(docx))
    xml = doc.element.xml
    problems = []
    if doc.tables:
        problems.append(f"{len(doc.tables)} table(s)")
    if "txbxContent" in xml:
        problems.append("text box")
    if "<w:drawing" in xml or "<w:pict" in xml:
        problems.append("image or drawing")
    columns = re.findall(r'<w:cols[^>]*w:num="(\d+)"', xml)
    if any(int(n) > 1 for n in columns):
        problems.append("multiple columns")
    for section in doc.sections:
        if any(p.text.strip() for p in section.header.paragraphs) or any(p.text.strip() for p in section.footer.paragraphs):
            problems.append("text in the header or footer")
            break
    body = textutil.squash("\n".join(p.text for p in doc.paragraphs))
    lost = [a for a in atoms if textutil.squash(a) and textutil.squash(a) not in body]
    if lost:
        problems.append(f"{len(lost)} piece(s) of text not in the document body")
    return AtsCheck(name="docx_structure", passed=not problems, score=0.0 if problems else 1.0,
                    detail="; ".join(problems) or "single column; no tables, text boxes, images, or header text")


# ---------------------------------------------------------------------------
# PDF helpers
# ---------------------------------------------------------------------------


def extract_pdf_text(pdf: Path) -> str:
    from pypdf import PdfReader

    return "\n".join((page.extract_text() or "") for page in PdfReader(str(pdf)).pages)


def pdf_fonts(pdf: Path) -> list[str]:
    """Names of the fonts embedded in the PDF (subset prefixes removed)."""
    from pypdf import PdfReader

    names: set[str] = set()
    for page in PdfReader(str(pdf)).pages:
        fonts = (page.get("/Resources") or {}).get("/Font") or {}
        for key in fonts:
            base = str(fonts[key].get_object().get("/BaseFont", "")).lstrip("/")
            names.add(base.split("+", 1)[-1])
    return sorted(n for n in names if n)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def _score(checks: list[AtsCheck]) -> float:
    weights = [(config.ATS_CHECK_WEIGHTS.get(c.name, 1.0), c.score if not c.passed else 1.0) for c in checks]
    total = sum(w for w, _ in weights)
    return round(10.0 * sum(w * s for w, s in weights) / total, 2) if total else 0.0


def repetition_terms(analysis: JobAnalysis | None, view: ResumeView) -> list[str]:
    terms: list[str] = []
    if analysis is not None:
        terms += analysis.keywords + [s.name for s in analysis.all_skills()] + analysis.tools
    skills = view.section("skills")
    if skills is not None:
        for line in skills.lines:
            terms += [item.strip() for item in line.split(":", 1)[-1].split(",") if item.strip()]
    return terms


def check_text(view: ResumeView, analysis: JobAnalysis | None = None) -> AtsReport:
    """Text-level checks only (no file needed). Page count and extraction come later."""
    checks = [
        check_contact(view),
        check_headings(view),
        check_odd_characters(view),
        check_links(view),
        check_repetition(view, repetition_terms(analysis, view)),
        check_chronology(view),
    ]
    return AtsReport(checks=checks, score=_score(checks), final=False)


def check_rendered(view: ResumeView, pdf: Path, docx: Path, analysis: JobAnalysis | None = None,
                   resolve_links: bool = False) -> AtsReport:
    """Every check, on the final PDF and DOCX."""
    from pypdf import PdfReader

    extracted = extract_pdf_text(pdf)
    atoms = view.atoms()
    extraction = compare_extraction(atoms, extracted)
    fonts = pdf_fonts(pdf)
    wanted = textutil.squash(config.BODY_FONT)
    if fonts and not any(wanted in textutil.squash(f) for f in fonts):
        extraction.detail += f"; note: the PDF uses {', '.join(fonts)} instead of {config.BODY_FONT}"
    checks = [
        check_contact(view, extracted),
        extraction,
        check_headings(view, extracted),
        check_odd_characters(view, extracted),
        check_links(view, resolve=resolve_links),
        check_repetition(view, repetition_terms(analysis, view)),
        check_chronology(view),
        check_docx(docx, atoms),
        check_page_count(len(PdfReader(str(pdf)).pages)),
    ]
    return AtsReport(checks=checks, score=_score(checks), final=True)
