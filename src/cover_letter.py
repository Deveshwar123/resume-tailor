"""Optional cover letter (--cover-letter), built on the same facts, validation, and approval.

Claude writes the letter sentence by sentence. Every sentence about me cites the
facts it rests on, and goes through the same evidence checks as a resume bullet
(`validate.check_bullet`): no number, skill, tool, or level of ownership that
the cited facts do not support. My education entry is also accepted as evidence
("a first-year Data Engineering student at Texas A&M").

    hard violation  -> one retry with the specific problem, then the sentence is removed
    soft flag       -> I approve, edit, or remove it (removed automatically in batch mode)
    company claims  -> must be grounded in the posting or official / posting-tier research

A cover letter is prose, and no validator catches everything: read it before sending.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Literal

from pydantic import BaseModel

import config
from src import render, textutil, ui, validate
from src.errors import LLMError
from src.models import CompanyBrief, MasterResume, Severity, Violation
from src.validate import ValidationContext

MAX_WORDS = 280

LETTER_USER = """Write a cover letter for this job, using the fact bank above.

- Three or four short paragraphs, at most {max_words} words in total. Plain and direct: no cliches, no flattery, no restating the whole resume.
- Open with the role and why this team's work fits what the candidate has actually done. Then connect the two or three most relevant facts to what the role needs. Close briefly.
- Every sentence about the candidate follows the same rules as a resume bullet: it states only what its cited facts support, keeps numbers exactly as written, names only skills and tools the cited facts name, and keeps ownership at or below the facts' level. The education entry in the fact bank may be stated without a fact id.
- Split the letter into sentences. For each one give:
  - `about`: "me" (a claim about the candidate), "company" (a statement about the company or role), or "other" (greeting-style or closing lines with no factual claim).
  - `fact_ids`: the facts a "me" sentence rests on (they may come from different experiences). Empty for "company" and "other".
- Statements about the company must come from the job posting or the company research given above.
- `greeting`: for example "Dear Hiring Manager,". `closing`: for example "Sincerely,"."""

FIX_USER = """These cover-letter sentences failed validation against the fact bank. Rewrite each so it states only what its cited facts support and no longer has the problem listed. Keep each `key` unchanged.

{requests}"""


class LetterSentence(BaseModel):
    text: str
    about: Literal["me", "company", "other"]
    fact_ids: list[str]


class LetterParagraph(BaseModel):
    sentences: list[LetterSentence]


class CoverLetterDraft(BaseModel):
    """Wire model: the letter as cited sentences."""

    greeting: str
    paragraphs: list[LetterParagraph]
    closing: str


class SentenceFix(BaseModel):
    key: str
    text: str
    fact_ids: list[str]


class SentenceFixBatch(BaseModel):
    fixes: list[SentenceFix]


@dataclass
class Sentence:
    text: str
    about: str
    fact_ids: list[str]
    paragraph: int
    flags: list[str] = field(default_factory=list)      # soft flags waiting for me
    decision: str = "auto"                              # auto | approved | edited | removed
    notes: list[str] = field(default_factory=list)

    @property
    def kept(self) -> bool:
        return self.decision != "removed"


@dataclass
class CoverLetter:
    greeting: str
    closing: str
    sentences: list[Sentence] = field(default_factory=list)
    violations: list[Violation] = field(default_factory=list)
    files: dict[str, Path] = field(default_factory=dict)
    pages: int = 0

    def paragraphs(self) -> list[str]:
        grouped: dict[int, list[str]] = {}
        for sentence in self.sentences:
            if sentence.kept:
                grouped.setdefault(sentence.paragraph, []).append(sentence.text)
        return [" ".join(parts) for _index, parts in sorted(grouped.items()) if parts]

    def word_count(self) -> int:
        return sum(textutil.word_count(p) for p in self.paragraphs())

    def pending(self) -> list[Sentence]:
        return [s for s in self.sentences if s.flags and s.decision == "auto"]


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

_FIRST_PERSON = re.compile(r"\b(?:I|I'm|I've|[Mm]y|[Mm]e|[Mm]yself|[Mm]ine)\b")


def _education_evidence(master: MasterResume) -> list[str]:
    lines: list[str] = []
    for edu in master.education:
        lines += [edu.school, edu.degree, edu.major, edu.grad, edu.location, *edu.honors, *edu.coursework, *edu.notes]
        if edu.gpa is not None:
            lines.append(f"GPA {edu.gpa:.1f} {edu.gpa:.2f} {config.GPA_SCALE:.1f}")
        if edu.class_year:
            lines.append(f"class of {edu.class_year}")
    return [line for line in lines if line and line.strip()]


def grounded(text: str, sources: list[str]) -> bool:
    """Most content words of `text` appear in the posting or trusted research."""
    wanted = textutil.content_words(text)
    if not wanted:
        return True
    have = set(textutil.words(" ".join(sources)))
    stems = {w[:6] for w in have}
    return sum(1 for w in wanted if w in have or w[:6] in stems) / len(wanted) >= 0.6


def check_sentence(text: str, about: str, fact_ids: list[str], ctx: ValidationContext,
                   company_sources: list[str]) -> list[Violation]:
    """Evidence checks for one sentence. Company sentences are checked for grounding instead."""
    if about == "company" and not _FIRST_PERSON.search(text):
        if grounded(text, company_sources):
            return []
        return [Violation(code="L1", severity=Severity.soft, token="company",
                          message="says something about the company that the posting and trusted research do not state")]
    found: list[Violation] = []
    facts = []
    for fact_id in dict.fromkeys(fact_ids):
        hit = ctx.facts.get(fact_id)
        if hit is None:
            found.append(Violation(code="A1", severity=Severity.hard, token=fact_id,
                                   message=f'cites fact "{fact_id}", which does not exist'))
        else:
            facts.append(hit)
    entry = facts[0][0] if facts else next(iter(ctx.master.entries()), None)
    if entry is None:
        return found
    evidence = validate.gather_evidence(entry, [fact for _entry, fact in facts], None)
    if facts:
        tiers = [config.OWNERSHIP_TIERS[MasterResume.effective_ownership(e, f).value] for e, f in facts]
        evidence.max_tier, evidence.min_tier = max(tiers), min(tiers)
    else:
        evidence.max_tier = evidence.min_tier = config.OWNERSHIP_TIERS["contributor"]
    # Names that may appear in any letter: my school and degree, the employers and titles of the
    # cited facts, and the company and role I am writing to.
    extra = _education_evidence(ctx.master)
    for cited_entry, _fact in facts:
        extra += [cited_entry.display_org, cited_entry.display_title, cited_entry.location,
                  *textutil.years_in(cited_entry.dates.start, cited_entry.dates.end)]
    if ctx.analysis is not None:
        extra += [ctx.analysis.company, ctx.analysis.role, ctx.analysis.location]
    evidence.texts = [*evidence.texts, *[e for e in extra if e]]
    for violation in validate.check_bullet(text, evidence, ctx, entry.id, None):
        if violation.code == "C3":
            continue                      # a letter does not have to repeat every metric
        if violation.code == "B3" and "drops the helper wording" in violation.message and not facts:
            continue
        found.append(violation)
    return found


def company_sources(job_text: str, brief: CompanyBrief | None) -> list[str]:
    sources = [job_text]
    if brief is not None:
        for items in brief.items(config.JUDGE_TIERS).values():
            sources += [item.text for item in items]
    return sources


def validate_letter(draft: CoverLetterDraft, ctx: ValidationContext, sources: list[str], retry_fn=None) -> CoverLetter:
    """Check every sentence; retry hard failures once; remove what still fails."""
    letter = CoverLetter(greeting=textutil.plain(draft.greeting) or "Dear Hiring Manager,",
                         closing=textutil.plain(draft.closing) or "Sincerely,")
    checked: list[tuple[Sentence, list[Violation]]] = []
    for index, paragraph in enumerate(draft.paragraphs):
        for item in paragraph.sentences:
            text = textutil.plain(item.text)
            if not text:
                continue
            sentence = Sentence(text=text, about=item.about, fact_ids=list(item.fact_ids), paragraph=index)
            checked.append((sentence, check_sentence(text, item.about, item.fact_ids, ctx, sources)))

    failing = [(i, s, v) for i, (s, v) in enumerate(checked) if any(x.severity == Severity.hard for x in v)]
    fixes: dict[str, SentenceFix] = {}
    if failing and retry_fn is not None:
        fixes = {f.key: f for f in retry_fn([
            {"key": str(i), "sentence": s.text, "fact_ids": s.fact_ids,
             "problems": [x.message for x in v if x.severity == Severity.hard]} for i, s, v in failing])}
    for i, sentence, violations in failing:
        hard = [x for x in violations if x.severity == Severity.hard]
        fix = fixes.get(str(i))
        if fix is not None and textutil.plain(fix.text):
            retry = check_sentence(textutil.plain(fix.text), sentence.about, fix.fact_ids, ctx, sources)
            if not any(x.severity == Severity.hard for x in retry):
                for x in hard:
                    x.resolution = "fixed by the one allowed retry"
                letter.violations += hard
                sentence.text, sentence.fact_ids = textutil.plain(fix.text), list(fix.fact_ids)
                sentence.notes.append("rewritten once after a validation failure")
                checked[i] = (sentence, retry)
                continue
        for x in hard:
            x.resolution = "sentence removed" if retry_fn is None else "sent back to Claude once; sentence removed"
        letter.violations += hard
        sentence.decision = "removed"
        sentence.notes.append("removed: " + "; ".join(x.message for x in hard))
        checked[i] = (sentence, [])

    for sentence, violations in checked:
        soft = [x for x in violations if x.severity == Severity.soft]
        if sentence.kept and soft:
            sentence.flags = [x.message for x in soft]
            for x in soft:
                x.resolution = "sent to approval"
            letter.violations += soft
        letter.sentences.append(sentence)
    return letter


# ---------------------------------------------------------------------------
# Writing, approval, rendering
# ---------------------------------------------------------------------------


def write(llm, system: str, ctx: ValidationContext, job_text: str, brief: CompanyBrief | None) -> CoverLetter:
    """Ask Claude for the letter (same cached fact-bank prompt as the resume) and validate it."""
    draft = llm.parse(purpose="cover_letter", model=config.WRITE_MODEL, system=system,
                      user=LETTER_USER.format(max_words=MAX_WORDS), schema=CoverLetterDraft)

    def retry(requests: list[dict]) -> list[SentenceFix]:
        try:
            batch = llm.parse(purpose="fix", model=config.WRITE_MODEL, system=system,
                              user=FIX_USER.format(requests=json.dumps(requests, indent=1, ensure_ascii=False)),
                              schema=SentenceFixBatch)
        except LLMError as exc:
            ui.warn(f"Cover-letter retry could not run ({exc.message}); removing those sentences.")
            return []
        return batch.fixes

    return validate_letter(draft, ctx, company_sources(job_text, brief), retry)


def review(letter: CoverLetter, ctx: ValidationContext, sources: list[str], *, batch: bool = False) -> None:
    """Flagged sentences: [A]pprove / [E]dit / [R]emove. With no one at the keyboard they are removed."""
    pending = letter.pending()
    if not pending:
        return
    if batch or not ui.is_interactive():
        for sentence in pending:
            sentence.decision = "removed"
            sentence.notes.append("removed: flagged and not approved (no one at the keyboard)")
        return
    ui.console.print(f"\n[bold]{len(pending)} cover-letter sentence(s) need your approval[/bold] (Enter = remove)")
    for number, sentence in enumerate(pending, 1):
        ui.console.print(f"\n[{number}/{len(pending)}] {sentence.text}")
        for flag in sentence.flags:
            ui.console.print(f"   [yellow]flagged:[/yellow] {flag}")
        if sentence.fact_ids:
            for fact_id in sentence.fact_ids:
                hit = ctx.facts.get(fact_id)
                if hit:
                    ui.console.print(f"   [dim]fact {fact_id}: {hit[1].text}[/dim]")
        choice = ui.ask_choice("  ", {"a": "approve", "e": "edit", "r": "remove"}, default="r")
        if choice == "a":
            sentence.decision = "approved"
        elif choice == "e":
            edited = textutil.plain(ui.ask("   Your wording (empty to remove): "))
            hard = [v for v in check_sentence(edited, sentence.about, sentence.fact_ids, ctx, sources)
                    if v.severity == Severity.hard] if edited else []
            if edited and not hard:
                sentence.text, sentence.decision = edited, "edited"
            else:
                for violation in hard:
                    ui.console.print(f"   [red]not supported:[/red] {violation.message}")
                sentence.decision = "removed"
        else:
            sentence.decision = "removed"
        sentence.notes.append(f"approval: {sentence.decision}")


def to_markdown(letter: CoverLetter, master: MasterResume, company: str, role: str, today: date) -> str:
    lines = [f"# Cover letter - {role} at {company}", "", f"{today:%B} {today.day}, {today.year}", "", letter.greeting, ""]
    for paragraph in letter.paragraphs():
        lines += [paragraph, ""]
    lines += [letter.closing, master.contact.name, "", "---", "", "## Where each sentence comes from", "",
              "| Sentence | About | Facts | Decision | Notes |", "|---|---|---|---|---|"]
    for sentence in letter.sentences:
        notes = "; ".join([*sentence.flags, *sentence.notes]).replace("|", "\\|")
        lines.append(f"| {sentence.text.replace('|', chr(92) + '|')} | {sentence.about} | "
                     f"{', '.join(sentence.fact_ids) or '-'} | {sentence.decision} | {notes} |")
    lines += ["", f"{letter.word_count()} words. A cover letter is prose: read it once before sending."]
    return "\n".join(lines) + "\n"


def render_docx(letter: CoverLetter, master: MasterResume, path: Path, today: date) -> Path:
    """A plain one-column letter in the resume's font."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1.0))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = config.BODY_FONT, Pt(config.DEFAULT_FONT_PT)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    fonts = normal.element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is not None:
        for attr in list(fonts.attrib):
            del fonts.attrib[attr]
        for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(slot), config.BODY_FONT)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.08

    contact = master.contact
    header = doc.add_paragraph()
    name = header.add_run(textutil.plain(contact.name))
    name.bold = True
    header.paragraph_format.space_after = Pt(0)
    parts = [contact.location, contact.phone, contact.email, contact.linkedin]
    line = doc.add_paragraph(render.CONTACT_SEPARATOR.join(textutil.plain(p) for p in parts if p and p.strip()))
    line.paragraph_format.space_after = Pt(14)
    doc.add_paragraph(f"{today:%B} {today.day}, {today.year}")
    doc.add_paragraph(letter.greeting)
    for paragraph in letter.paragraphs():
        doc.add_paragraph(paragraph)
    closing = doc.add_paragraph(letter.closing)
    closing.paragraph_format.space_after = Pt(0)
    doc.add_paragraph(textutil.plain(contact.name))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def save(letter: CoverLetter, master: MasterResume, out_dir: Path, company: str, role: str, today: date) -> CoverLetter:
    """Write cover_letter.md / .docx / .pdf into the run's output folder."""
    out_dir.mkdir(parents=True, exist_ok=True)
    markdown = out_dir / "cover_letter.md"
    markdown.write_text(to_markdown(letter, master, company, role, today), encoding="utf-8")
    letter.files["cover_letter_md"] = markdown
    docx = render_docx(letter, master, out_dir / "cover_letter.docx", today)
    letter.files["cover_letter_docx"] = docx
    pdf = render.docx_to_pdf(docx, out_dir)
    letter.files["cover_letter_pdf"] = pdf
    letter.pages = render.page_count(pdf)
    return letter
