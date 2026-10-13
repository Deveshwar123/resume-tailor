#!/usr/bin/env python3
"""Build template.docx - the ATS-safe resume design that render.py fills in.

Run this to regenerate the template from scratch:

    python scripts/build_template.py

You can also edit template.docx directly in Word (fonts, sizes, spacing live in
the paragraph styles named "Resume ..."). Keep it ATS-safe: one column, no
tables or text boxes, plain section headings, a common font, and leave the
{{ ... }} and {%p ... %} placeholders exactly as they are.

Design
------
* Single column, no tables, no text boxes, no headers/footers, no images.
* One common font throughout (config.BODY_FONT), black text.
* Section headings are real paragraphs with a bottom rule (a paragraph border,
  not a drawing) and an outline level, so parsers and screen readers see them.
* Dates sit on the title line at a right tab stop (plain text, not a table).
* Bullets are a literal "•" plus a tab with a hanging indent, so the text that
  comes out of the PDF is exactly the text that went in.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config  # noqa: E402

PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
BULLET_INDENT_IN = 0.22


def _set_fonts(rpr_parent, font: str) -> None:
    """Set every font slot explicitly and remove theme fonts so the choice sticks."""
    rpr = rpr_parent.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_parent.append(rpr)
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in list(fonts.attrib):
        del fonts.attrib[attr]
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(slot), font)


def _bottom_rule(style) -> None:
    ppr = style.element.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    ppr.append(borders)


def _outline_level(style, level: int) -> None:
    ppr = style.element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    ppr.append(outline)


def _style(doc, name: str, *, size: float, bold: bool = False, align=None, before: float = 0, after: float = 0,
           keep_next: bool = False, left_in: float = 0, hanging_in: float = 0):
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.quick_style = True
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    _set_fonts(style.element, config.BODY_FONT)
    fmt = style.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = 1.0
    fmt.keep_with_next = keep_next
    fmt.widow_control = True
    if align is not None:
        fmt.alignment = align
    if left_in:
        fmt.left_indent = Inches(left_in)
        fmt.first_line_indent = Inches(-hanging_in)
    return style


def build(path: Path) -> Path:
    doc = Document()

    # Page: US Letter, equal margins.
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(PAGE_WIDTH_IN), Inches(PAGE_HEIGHT_IN)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(config.DEFAULT_MARGIN_IN))
    text_width = PAGE_WIDTH_IN - 2 * config.DEFAULT_MARGIN_IN

    # Base font: set on the document defaults and on Normal, with theme fonts removed.
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None:
            _set_fonts(rpr_default, config.BODY_FONT)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(config.DEFAULT_FONT_PT)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    _set_fonts(normal.element, config.BODY_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    body = config.DEFAULT_FONT_PT
    name = _style(doc, "Resume Name", size=body + 5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    contact = _style(doc, "Resume Contact", size=body, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    heading = _style(doc, "Resume Heading", size=body + 1, bold=True, before=7, after=2, keep_next=True)
    _bottom_rule(heading)
    _outline_level(heading, 0)
    entry = _style(doc, "Resume Entry", size=body, bold=True, before=3, keep_next=True)
    entry.paragraph_format.tab_stops.add_tab_stop(Inches(text_width), WD_TAB_ALIGNMENT.RIGHT)
    line = _style(doc, "Resume Line", size=body)
    bullet = _style(doc, "Resume Bullet", size=body, before=1, left_in=BULLET_INDENT_IN, hanging_in=BULLET_INDENT_IN)
    bullet.paragraph_format.tab_stops.add_tab_stop(Inches(BULLET_INDENT_IN))

    def add(text: str, style) -> None:
        doc.add_paragraph(text, style=style)

    add("{{ name }}", name)
    add("{{ contact_line }}", contact)
    add("{%p for section in sections %}", line)
    add("{{ section.heading }}", heading)
    add("{%p for entry in section.entries %}", line)
    add("{{ entry.left }}\t{{ entry.right }}", entry)
    add("{%p for text in entry.lines %}", line)
    add("{{ text }}", line)
    add("{%p endfor %}", line)
    add("{%p for text in entry.bullets %}", line)
    add("•\t{{ text }}", bullet)
    add("{%p endfor %}", line)
    add("{%p endfor %}", line)
    add("{%p for text in section.lines %}", line)
    add("{{ text }}", line)
    add("{%p endfor %}", line)
    add("{%p for text in section.bullets %}", line)
    add("•\t{{ text }}", bullet)
    add("{%p endfor %}", line)
    add("{%p endfor %}", line)

    doc.core_properties.title = "Resume"
    doc.core_properties.author = ""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else config.TEMPLATE_PATH
    print(f"wrote {build(target)}")
